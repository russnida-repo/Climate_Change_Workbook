#!/usr/bin/env python3
"""
make_figures.py

Draws the five landing-page figures from Climate_Change_Analysis_2.xlsx.

The figures are generated from the workbook's own cells rather than exported
from Excel, for the same reason the four validation scripts exist: a figure
nobody can regenerate is a figure nobody can check. Every series below names
the cells it reads, so a reviewer can verify any curve against the workbook or
against the Markdown export.

Each figure is written twice, once for a light surface and once for a dark one.
The dark variants are separately chosen steps from the same hues, not an
automatic inversion. The page pairs them with <picture> and a
prefers-color-scheme media query.

Text is baked to vector paths (svg.fonttype = "path") so each file renders
identically wherever it is embedded, including in a GitHub README, with no
webfont dependency.

Palette provenance: the three categorical hues and the ordinal warm ramp were
checked with the data-visualisation validator against this page's own surfaces
(#fbfcfd light, #0f141a dark) on the all-pairs list. Every check passes in both
modes: lightness band, chroma floor, colour-vision-deficiency separation,
normal-vision separation and contrast against the surface.

Usage
-----
  python3 make_figures.py [workbook.xlsx] [output_dir]

Defaults: Climate_Change_Analysis_2.xlsx and docs/figures.

Dependencies: matplotlib, openpyxl.
"""

import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.path import Path
from matplotlib.patches import PathPatch
import openpyxl

matplotlib.rcParams["svg.fonttype"] = "path"
matplotlib.rcParams["font.family"] = [
    "Public Sans", "Inter", "DejaVu Sans", "sans-serif"]

DEFAULT_WB = "Climate_Change_Analysis_2.xlsx"
DEFAULT_OUT = os.path.join("docs", "figures")

# ----------------------------------------------------------------------------
# Themes. Dark is selected for the dark surface, not flipped.
# ----------------------------------------------------------------------------
THEMES = {
    "light": dict(
        surface="#fbfcfd",
        ink="#10161b", ink2="#4a5560", ink3="#6b7680",
        grid="#e3e8ec", axis="#c6ced5",
        blue="#1a5fb4", orange="#d4541f", teal="#0f8a62",
        ramp=["#9c3a14", "#d4541f", "#e8894f"],   # ordinal, dark -> light
        cool="#1a5fb4",
        open="#a5272c",
    ),
    "dark": dict(
        surface="#0f141a",
        ink="#e6ecf1", ink2="#a3b0bc", ink3="#7d8b98",
        grid="#232c35", axis="#39454f",
        blue="#4a93ea", orange="#e0662f", teal="#1fa877",
        ramp=["#a8431a", "#e0662f", "#f0a476"],
        cool="#4a93ea",
        open="#f08a8f",
    ),
}


def isnum(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def col(ws, letter_or_idx, r0, r1):
    """Read a column over a row range, keeping None for blanks."""
    if isinstance(letter_or_idx, str):
        from openpyxl.utils import column_index_from_string
        c = column_index_from_string(letter_or_idx)
    else:
        c = letter_or_idx
    return [ws.cell(r, c).value for r in range(r0, r1 + 1)]


def frame(ax, t, xlabel=None, ylabel=None, xgrid=False):
    """Recessive chrome: hairline grid behind, two thin spines, muted ticks."""
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color(t["axis"])
        ax.spines[side].set_linewidth(0.8)
    ax.tick_params(colors=t["ink3"], labelsize=9, length=3, width=0.8)
    ax.set_axisbelow(True)
    ax.grid(axis="x" if xgrid else "y", color=t["grid"],
            linewidth=0.8, alpha=1.0)
    if xlabel:
        ax.set_xlabel(xlabel, color=t["ink2"], fontsize=9.5, labelpad=7)
    if ylabel:
        ax.set_ylabel(ylabel, color=t["ink2"], fontsize=9.5, labelpad=7)


def titles(fig, t, title, subtitle, x=0.012, y=0.975):
    fig.text(x, y, title, color=t["ink"], fontsize=13.5,
             fontweight="600", va="top", ha="left")
    if subtitle:
        fig.text(x, y - 0.085, subtitle, color=t["ink2"], fontsize=10,
                 va="top", ha="left")


def source(fig, t, text, y=0.015):
    fig.text(0.012, y, text, color=t["ink3"], fontsize=8.4,
             va="bottom", ha="left")


def _px_to_data(ax, px):
    """Return (dx, dy): how many data units `px` screen pixels span on each axis.

    The two axes here are on completely different scales (°C against a category
    index), so a single radius in data units would round one axis correctly and
    deform the other. Call this only after the limits are set.
    """
    inv = ax.transData.inverted()
    x0, y0 = inv.transform((0, 0))
    x1, y1 = inv.transform((px, px))
    return abs(x1 - x0), abs(y1 - y0)


def rounded_barh(ax, y, width, height, color, px=4, zorder=2):
    """Horizontal bar with a rounded data end and a square baseline end."""
    rx, ry = _px_to_data(ax, px)
    rx = min(rx, abs(width) * 0.5)
    ry = min(ry, height * 0.5)
    s = 1.0 if width >= 0 else -1.0
    x0, x1 = 0.0, width
    y0, y1 = y - height / 2, y + height / 2
    xin = x1 - s * rx
    verts = [(x0, y0), (xin, y0), (x1, y0), (x1, y0 + ry),
             (x1, y1 - ry), (x1, y1), (xin, y1), (x0, y1), (x0, y0)]
    codes = [Path.MOVETO, Path.LINETO, Path.CURVE3, Path.CURVE3,
             Path.LINETO, Path.CURVE3, Path.CURVE3, Path.LINETO,
             Path.CLOSEPOLY]
    ax.add_patch(PathPatch(Path(verts, codes), facecolor=color,
                           edgecolor="none", zorder=zorder))


def newfig(t, w=8.6, h=4.9):
    fig, ax = plt.subplots(figsize=(w, h))
    fig.patch.set_facecolor(t["surface"])
    ax.set_facecolor(t["surface"])
    return fig, ax


def save(fig, out_dir, name, mode, t):
    os.makedirs(out_dir, exist_ok=True)
    p = os.path.join(out_dir, f"{name}-{mode}.svg")
    fig.savefig(p, format="svg", facecolor=t["surface"],
                edgecolor="none", bbox_inches="tight", pad_inches=0.18)
    plt.close(fig)
    return p


# ----------------------------------------------------------------------------
# Figure 1. What each control buys.   Summary!M12:O15
# ----------------------------------------------------------------------------
def fig_controls(wb, t, out_dir, mode):
    ws = wb["Summary"]
    labels, r2, sd = [], [], []
    for r in range(12, 16):
        labels.append(str(ws.cell(r, 13).value))
        r2.append(float(ws.cell(r, 14).value))
        sd.append(float(ws.cell(r, 15).value))

    pretty = {"CO2 only": "CO₂ only", "+ ENSO": "plus ENSO",
              "+ volcanic": "plus volcanic aerosol",
              "kernels": "exponential responses"}
    labels = [pretty.get(l, l) for l in labels]

    fig, ax = newfig(t, h=4.3)
    ypos = list(range(len(labels)))[::-1]
    # limits before the bars: the corner radius is computed in pixels
    ax.set_xlim(0, max(sd) * 1.25)
    ax.set_ylim(-0.75, len(labels) - 0.25)
    for y, v in zip(ypos, sd):
        rounded_barh(ax, y, v, 0.52, t["blue"])
    for y, v, q in zip(ypos, sd, r2):
        ax.text(v + 0.0032, y, f"{v:.3f}", va="center", ha="left",
                color=t["ink"], fontsize=10, fontweight="600")
        ax.text(v + 0.0032, y - 0.30, f"R² {q:.3f}", va="center", ha="left",
                color=t["ink3"], fontsize=8.8)
    ax.set_yticks(ypos)
    ax.set_yticklabels(labels, fontsize=10, color=t["ink"])
    frame(ax, t, xlabel="Residual standard deviation (°C)", xgrid=True)
    ax.grid(axis="y", visible=False)
    titles(fig, t, "Each control removes real variance",
           "Monthly fits, 1958 to 2026. Shorter is better. The scale starts at zero,\n"
           "so the bars are comparable by length.")
    source(fig, t, "Summary!M12:O15")
    fig.subplots_adjust(top=0.74, bottom=0.17, left=0.21, right=0.97)
    return save(fig, out_dir, "fig1-controls", mode, t)


# ----------------------------------------------------------------------------
# Figure 2. The headline regression.   Forcing!A,B,AN rows 2:68
# ----------------------------------------------------------------------------
def fig_forcing_fit(wb, t, out_dir, mode):
    F = wb["Forcing"]
    yr = col(F, "A", 2, 68)
    an = col(F, "B", 2, 68)
    er = col(F, "AN", 2, 68)
    pts = [(float(a), float(b), int(y)) for y, b, a in zip(yr, an, er)
           if isnum(y) and isnum(a) and isnum(b)]
    xs = [p[0] for p in pts]
    ys = [p[1] for p in pts]

    b = float(F["T14"].value)
    a0 = float(F["T13"].value)
    c = float(F["T15"].value)
    oni = [float(v) for v in col(F, "AO", 2, 68) if isnum(v)]
    mean_oni = sum(oni) / len(oni)

    fig, ax = newfig(t)
    x0, x1 = min(xs) - 0.12, max(xs) + 0.12
    ax.plot([x0, x1],
            [a0 + b * x0 + c * mean_oni, a0 + b * x1 + c * mean_oni],
            color=t["ink2"], linewidth=2.0, zorder=3, solid_capstyle="round")
    ax.scatter(xs, ys, s=34, color=t["blue"], alpha=0.85,
               edgecolors=t["surface"], linewidths=1.0, zorder=4)

    # The slope sits in the empty upper-left rather than beside the line, where
    # it collided with both the points and the end-year labels.
    ax.text(0.025, 0.95, f"slope {b:.4f} °C per W/m²",
            transform=ax.transAxes, va="top", ha="left",
            fontsize=9.6, color=t["ink2"], fontweight="500")

    ax.set_xlim(x0, x1)
    frame(ax, t, xlabel="Anthropogenic effective radiative forcing (W/m²)",
          ylabel="Annual mean anomaly (°C)")
    titles(fig, t, "Temperature against human forcing, 1959 to 2025",
           "67 annual means, ENSO controlled. R² 0.953. The line is the fitted\n"
           "regression at mean ENSO, not a trend through the points.")
    source(fig, t, "Forcing!B2:B68 against Forcing!AN2:AN68; fit Forcing!T13:T17")
    fig.subplots_adjust(top=0.76, bottom=0.14, left=0.085, right=0.975)
    return save(fig, out_dir, "fig2-forcing-fit", mode, t)


# ----------------------------------------------------------------------------
# Figure 3. Forcing by component.   Forcing!K, I, L rows 2:68
# ----------------------------------------------------------------------------
def fig_components(wb, t, out_dir, mode):
    F = wb["Forcing"]
    yr = [int(v) for v in col(F, "A", 2, 68) if isnum(v)]
    ghg = [float(v) for v in col(F, "K", 2, 68) if isnum(v)]
    aer = [float(v) for v in col(F, "I", 2, 68) if isnum(v)]
    tot = [float(v) for v in col(F, "L", 2, 68) if isnum(v)]

    fig, ax = newfig(t)
    ax.axhline(0, color=t["axis"], linewidth=0.9, zorder=1)
    series = [("Greenhouse gases", ghg, t["orange"]),
              ("Anthropogenic total", tot, t["blue"]),
              ("Aerosols", aer, t["teal"])]
    for name, vals, colr in series:
        ax.plot(yr, vals, color=colr, linewidth=2.0, zorder=3,
                solid_capstyle="round", label=name)
    # Direct labels sit in the right margin, so identity never rests on colour
    # alone and the axis still stops where the data stops.
    for name, vals, colr in series:
        ax.annotate(name, (yr[-1], vals[-1]), textcoords="offset points",
                    xytext=(10, 0), va="center", fontsize=9.6,
                    color=t["ink2"], fontweight="500", annotation_clip=False)

    ax.set_xlim(yr[0], yr[-1])
    frame(ax, t, ylabel="Effective radiative forcing (W/m²)")
    titles(fig, t, "The aerosol series is what makes the split identifiable",
           "Greenhouse-gas forcing trends while aerosol forcing is U-shaped, so the two\n"
           "are nearly orthogonal over the record (r = −0.002).")
    source(fig, t, "Forcing!K, I and L, rows 2:68")
    fig.subplots_adjust(top=0.76, bottom=0.11, left=0.085, right=0.80)
    return save(fig, out_dir, "fig3-erf-components", mode, t)


# ----------------------------------------------------------------------------
# Figure 4. Attribution.   Forcing!BB:BE row 68, with BF observed
# ----------------------------------------------------------------------------
def fig_attribution(wb, t, out_dir, mode):
    F = wb["Forcing"]
    comps = [("CO₂", float(F["BB68"].value), t["ramp"][0]),
             ("Other greenhouse gases", float(F["BC68"].value), t["ramp"][1]),
             ("Ozone, land use, other", float(F["BD68"].value), t["ramp"][2]),
             ("Aerosols", float(F["BE68"].value), t["cool"])]
    fitted = float(F["T28"].value)
    observed = float(F["BF68"].value)

    fig, ax = newfig(t, h=4.6)
    ypos = list(range(len(comps)))[::-1]
    # limits before the bars: the corner radius is computed in pixels
    ax.set_xlim(-0.18, max(c[1] for c in comps) * 1.30)
    ax.set_ylim(-0.95, len(comps) - 0.3)
    for y, (name, v, colr) in zip(ypos, comps):
        rounded_barh(ax, y, v, 0.5, colr)
        off = 0.012 if v >= 0 else -0.012
        ax.text(v + off, y, f"{v:+.3f}", va="center",
                ha="left" if v >= 0 else "right",
                color=t["ink"], fontsize=10, fontweight="600")
    ax.axvline(0, color=t["axis"], linewidth=0.9, zorder=1)
    ax.set_yticks(ypos)
    ax.set_yticklabels([c[0] for c in comps], fontsize=10, color=t["ink"])
    frame(ax, t, xlabel="Fitted warming since 1959 (°C)", xgrid=True)
    ax.grid(axis="y", visible=False)
    # the zero line is the reference here, so a left spine only competes with it
    ax.spines["left"].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.text(0, -0.86, f"Sum of the four: {fitted:.3f} °C      "
                      f"Observed over the same span: {observed:.3f} °C",
            fontsize=9.6, color=t["ink2"], va="center", ha="left")
    titles(fig, t, "Where the fitted warming came from",
           "Component contributions at 2025, from the anthropogenic-forcing fit.\n"
           "Aerosols are the one negative term.")
    source(fig, t, "Forcing!BB68:BE68, BF68 and T28")
    fig.subplots_adjust(top=0.75, bottom=0.17, left=0.215, right=0.97)
    return save(fig, out_dir, "fig4-attribution", mode, t)


# ----------------------------------------------------------------------------
# Figure 5. What is left over.   Kernel!B, M rows 9:815
# ----------------------------------------------------------------------------
def fig_residual(wb, t, out_dir, mode):
    K = wb["Kernel"]
    tt = col(K, "B", 9, 815)
    rr = col(K, "M", 9, 815)
    pts = [(float(a), float(b)) for a, b in zip(tt, rr)
           if isnum(a) and isnum(b)]
    xs = [p[0] for p in pts]
    ys = [p[1] for p in pts]

    fig, ax = newfig(t)
    ax.axhline(0, color=t["axis"], linewidth=0.9, zorder=2)
    ax.plot(xs, ys, color=t["blue"], linewidth=1.0, alpha=0.9, zorder=3)
    late = [(x, y) for x, y in pts if x >= 2023.0]
    if late:
        ax.plot([p[0] for p in late], [p[1] for p in late],
                color=t["open"], linewidth=1.6, zorder=4)
        mean_late = sum(p[1] for p in late) / len(late)
        ax.annotate(f"2023 to 2025 mean {mean_late:+.2f} °C",
                    (late[len(late) // 2][0], max(p[1] for p in late)),
                    textcoords="offset points", xytext=(-6, 16),
                    ha="right", fontsize=9.4, color=t["open"],
                    fontweight="500")
    frame(ax, t, ylabel="Residual (°C)")
    ax.set_xlim(min(xs), max(xs) + 0.4)
    titles(fig, t, "What the model does not explain",
           "Monthly residual from the exponential-response model. The recent excess is\n"
           "stated as an open item rather than attributed.")
    source(fig, t, "Kernel!M9:M815 against Kernel!B9:B815")
    fig.subplots_adjust(top=0.76, bottom=0.11, left=0.085, right=0.975)
    return save(fig, out_dir, "fig5-residual", mode, t)


FIGURES = [fig_controls, fig_forcing_fit, fig_components,
           fig_attribution, fig_residual]


def main(wb_path=DEFAULT_WB, out_dir=DEFAULT_OUT):
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    ver = wb["Summary"]["C48"].value
    print(f"workbook {wb_path} (version {ver})")
    written = []
    for mode, t in THEMES.items():
        for fn in FIGURES:
            written.append(fn(wb, t, out_dir, mode))
    for p in written:
        print(f"  wrote {p}  ({os.path.getsize(p):,} bytes)")
    print(f"{len(written)} files in {out_dir}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_WB,
         sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUT)
