#!/usr/bin/env python3
"""
open_items.py

The three questions left open by the September 2026 review and by the
2026-10-22 review round, worked through against Climate_Change_Analysis_2.xlsx:

  A. RONI instead of ONI as the ENSO control (flagged untested since September).
  B. Whether the GHG-versus-aerosol response split is identified, and if not,
     exactly which quantity is unidentified.
  C. Whether the CERES reflected-shortwave decline can account for the ~0.07 degC
     that remains in 2023-2025 after solar.

It computes every figure on the workbook's "Open Items" tab, and checks the
baselines it starts from against the workbook so a reader can see the comparison
is like for like.

Data added for this work
------------------------
RONI v5 (Relative Oceanic Nino Index, ERSSTv5 basis, 1991-2020 base period),
NOAA CPC, https://www.cpc.ncep.noaa.gov/products/analysis_monitoring/enso/roni/v5/
retrieved 2026-10-22. CPC publishes this series as an HTML table only, with no
ascii file, so it was transcribed into roni_v5_cpc.csv; that file carries a note
saying to re-check it against the page before publication. Everything else comes
from the workbook.

Findings, in brief
------------------
A. REPORTED AS A SENSITIVITY, not closed (revised v1.6). Swapping RONI for ONI
   makes the fit worse, not better (monthly R^2 0.9008 ->
   0.8918) and raises the CO2 coefficient from 2.727 to 2.820. All of that shift
   is RONI's own trend: RONI lag 3 falls 0.083 degC per decade over the sample
   while ONI lag 3 is flat at 0.005, and after removing RONI's linear trend the
   coefficient returns to 2.727, matching the ONI baseline to 0.0001. So the
   coefficient shift is entirely attributable to RONI's secular trend over this
   sample; that much is demonstrated. On where the trend comes from, two
   separate things can be said, and they are not the same kind of claim. By
   construction RONI is the Nino 3.4 anomaly minus the tropical-mean SST
   anomaly, so it declines mechanically if the tropical mean warms faster than
   Nino 3.4. Empirically the divergence tracks the warming: RONI minus ONI
   correlates -0.68 with the GISTEMP anomaly, and the anomaly explains variance
   in that divergence beyond a linear time trend (R^2 0.34 -> 0.48). Neither
   establishes causation inside this regression. What follows for the choice of
   control does not depend on settling that: a regressor carrying a secular
   trend competes with log2(CO2) for the same variance, so the trendless ONI is
   the appropriate control here. Two cautions on how far that argument goes.
   Detrending any regressor hands its trend variance back to log2(CO2), which is
   nearly linear in time, so the detrending result is close to a tautology and is
   not evidence about which index is right; and ONI's own flat trend is equally a
   product of its convention, its re-centred 30-year bases. The record cannot
   settle whether a forced change in the tropical Pacific SST gradient belongs
   with the ENSO term, so the honest statement is that the ENSO-index choice is a
   structural uncertainty of about +0.06 degC per doubling on the headline
   (1.607 under ONI, 1.665 under RONI), roughly twice the dataset sensitivity. The 2023-2025 residual gets larger under RONI,
   not smaller, because RONI makes the 2023-24 El Nino weaker (peak 1.5 against
   ONI's 2.0), leaving more of the recent warmth unexplained.

B. CORRECTED 2026-10-22 (v1.6). Until that revision this split regression used
   GHG ERF alone as the warming regressor, omitting Forcing column J, the other
   anthropogenic terms (ozone, land use, contrails, BC on snow, stratospheric
   water vapour), even though the headline regressor is L = K + I + J. Column J
   trends 0.13 to 0.48 W/m2 and correlates 0.98 with GHG, so its effect was
   absorbed into g: the reported 0.467 against 0.367, a difference of 0.101, was
   an omitted-variable artefact. With J restored the two responses are 0.409 and
   0.414, a difference of -0.005 (t = -0.05), i.e. indistinguishable, which is
   what the headline single-ERF regression assumes.

   What survives the correction is the identification result. The warming-term
   and aerosol series are nearly orthogonal over 1959-2025, because aerosol
   forcing is U-shaped while the warming terms trend, so the regression
   separates them and the two estimates are almost uncorrelated. What it cannot
   do is pin the aerosol response per W/m2, because the aerosol series magnitude
   is itself uncertain by a factor of four at AR6's 5-95%. Scaling that series by
   s scales h by exactly 1/s and leaves h*s and the R^2 unchanged, so the
   identified quantity is the aerosol contribution, not the response per unit
   forcing. Note also what a failure to reject does not buy: the 95% interval on
   g - h is -0.20 to +0.19, so the test is weak, and the right reading is that
   the data give no reason to prefer different responses.

C. The CERES shortwave decline cannot be treated as an independent forcing
   line. Reflected SW for 2023-2025 is 1.29 W/m2 below the 2001-2022 mean, which
   at the fitted response would be +0.53 degC: about 7.7 times the remainder it
   is meant to explain, and 1.50 times the whole change in total ERF across the
   same two windows. (Against the change in the CERES net imbalance rather than
   in total ERF the ratio is 2.38; the two denominators are different and should
   not be conflated.) A TOA flux change mixes externally forced cloud and albedo
   change, feedback to warming, and internal variability, and these numbers do
   not decompose it: what they establish is that the whole SW change cannot be
   entered as forcing, since the implied warming exceeds both the residual and
   the independently estimated forcing change, and since the regression's
   response coefficient already contains whatever part of it is feedback. Over
   the same window outgoing longwave rose 1.03 W/m2 while the net imbalance rose
   only 0.54. The workbook was right to leave it out of the residual budget. The
   ~0.07 degC remains unexplained.

Usage
-----
  python3 open_items.py <workbook.xlsx> <roni_v5_cpc.csv>

Dependencies: numpy, openpyxl.
"""

import math
import sys

import numpy as np
import openpyxl

F2X = 3.93
ROW_LO, ROW_HI = 9, 822          # the workbook's monthly fit window
ANNUAL_ROWS = (2, 69)            # Forcing rows 2..68 = 1959-2025


def isnum(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def read_roni(path):
    """{(year, month): value}; the 12 columns are seasons centred on Jan..Dec."""
    out = {}
    with open(path, encoding="utf-8-sig") as fh:
        for line in fh:
            if line.startswith("#") or line.startswith("year"):
                continue
            p = line.rstrip("\n").split(",")
            if not p or not p[0]:
                continue
            y = int(p[0])
            for m, cell in enumerate(p[1:13], start=1):
                if cell:
                    out[(y, m)] = float(cell)
    return out


def lagged(d, y, m, k):
    mm, yy = m - k, y
    while mm <= 0:
        mm += 12
        yy -= 1
    return d.get((yy, mm))


def ols(y, cols):
    A = np.column_stack([np.ones(len(y))] + [np.asarray(c, float) for c in cols])
    beta, *_ = np.linalg.lstsq(A, y, rcond=None)
    e = y - A @ beta
    n, k = A.shape
    s2 = (e ** 2).sum() / (n - k)
    V = np.linalg.inv(A.T @ A) * s2
    rho = np.corrcoef(e[:-1], e[1:])[0, 1]
    r2 = 1.0 - (e ** 2).sum() / ((y - y.mean()) ** 2).sum()
    return dict(b=beta, V=V, se=np.sqrt(np.diag(V)), r2=r2, rho=rho,
                infl=math.sqrt((1 + rho) / (1 - rho)), e=e, n=n,
                sd=e.std(ddof=1))


def two_sided_p(t):
    return 2 * (1 - 0.5 * (1 + math.erf(abs(t) / math.sqrt(2))))


def main(wb_path, roni_path):
    roni = read_roni(roni_path)
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    V, E, F, C = wb["Volcanic"], wb["ENSO"], wb["Forcing"], wb["CERES"]

    print(f"Workbook: {wb_path}")
    print(f"RONI:     {roni_path}  ({len(roni)} months, "
          f"{min(roni)} to {max(roni)})\n")

    # ---------------- monthly frame on the workbook's fit window
    rows = []
    for r in range(ROW_LO, ROW_HI + 1):
        a = V.cell(r, 1).value
        if not (isinstance(a, str) and "/" in a):
            continue
        m, y = (int(x) for x in a.split("/"))
        rows.append(dict(
            y=y, m=m, t=y + (m - 0.5) / 12,
            anom=V.cell(r, 2).value, lg=V.cell(r, 3).value,
            oni3=V.cell(r, 4).value, aod7=V.cell(r, 5).value,
            oni0=E.cell(r, 6).value, roni0=roni.get((y, m)),
            roni3=lagged(roni, y, m, 3)))
    sel = [d for d in rows
           if all(isnum(d[k]) for k in ("anom", "lg", "oni3", "aod7", "roni3"))]
    t = np.array([d["t"] for d in sel])
    for key in ("roni3", "oni3"):
        s = np.array([d[key] for d in sel], float)
        slope, inter = np.polyfit(t, s, 1)
        for d in sel:
            d[key + "_dt"] = d[key] - (slope * d["t"] + inter)
        print(f"   {key} linear trend over the fit window: "
              f"{slope * 10:+.4f} degC per decade")
    ov = [(d["roni0"], d["oni0"]) for d in rows
          if isnum(d.get("roni0")) and isnum(d.get("oni0"))]
    a1 = np.array([x[0] for x in ov]); a2 = np.array([x[1] for x in ov])
    print(f"   corr(RONI, ONI) on {len(ov)} same-month pairs: "
          f"{np.corrcoef(a1, a2)[0, 1]:.4f};  sd {a1.std(ddof=1):.3f} vs "
          f"{a2.std(ddof=1):.3f}")
    div = np.array([d["roni0"] - d["oni0"] for d in rows
                    if isnum(d.get("roni0")) and isnum(d.get("oni0"))
                    and isnum(d["anom"])])
    dan = np.array([d["anom"] for d in rows
                    if isnum(d.get("roni0")) and isnum(d.get("oni0"))
                    and isnum(d["anom"])], float)
    dt = np.array([d["t"] for d in rows
                   if isnum(d.get("roni0")) and isnum(d.get("oni0"))
                   and isnum(d["anom"])])
    f_t = ols(div, [dt])
    f_ta = ols(div, [dt, dan])
    print(f"   divergence RONI - ONI: corr with the GISTEMP anomaly "
          f"{np.corrcoef(div, dan)[0, 1]:+.4f}, slope "
          f"{np.polyfit(dan, div, 1)[0]:+.4f} per degC;"
          f" R^2 on time alone {f_t['r2']:.4f}, adding the anomaly "
          f"{f_ta['r2']:.4f} (anomaly coefficient {f_ta['b'][2]:+.4f})")
    print(f"   2023-24 peak: RONI {max(roni[(2023, m)] for m in range(10, 13)):+.1f}"
          f" against ONI {max(d['oni0'] for d in rows if d['y'] == 2023 and isnum(d['oni0'])):+.1f}")

    print("\nA. RONI instead of ONI, monthly three-factor model "
          "(identical months, n = %d)" % len(sel))
    print(f"   {'ENSO control':<24}{'b degC/doubling':>17}{'AR(1) SE':>10}"
          f"{'c':>8}{'R^2':>9}{'resid sd':>10}{'2023-25 resid':>15}")
    y = np.array([d["anom"] for d in sel], float)
    base_b = None
    for lab, key in (("ONI v5 (baseline)", "oni3"), ("RONI v5", "roni3"),
                     ("RONI v5, detrended", "roni3_dt")):
        f = ols(y, [[d["lg"] for d in sel], [d[key] for d in sel],
                    [d["aod7"] for d in sel]])
        res23 = float(np.mean([ev for d, ev in zip(sel, f["e"]) if d["y"] >= 2023]))
        if base_b is None:
            base_b = f["b"][1]
        print(f"   {lab:<24}{f['b'][1]:>17.4f}{f['se'][1]*f['infl']:>10.4f}"
              f"{f['b'][2]:>8.4f}{f['r2']:>9.4f}{f['sd']:>10.4f}{res23:>15.4f}")
    print(f"   workbook baseline: Volcanic!N15 b = {V['N15'].value:.4f}, "
          f"N18 R^2 = {V['N18'].value:.4f}, N16 c = {V['N16'].value:.4f}")

    # ---------------- annual
    ann = {}
    for d in sel:
        ann.setdefault(d["y"], []).append(d)
    arows = []
    for r in range(*ANNUAL_ROWS):
        yv = F.cell(r, 1).value
        if not isnum(yv):
            continue
        g = ann.get(int(yv), [])
        rl = [x["roni3"] for x in g if isnum(x["roni3"])]
        arows.append(dict(y=int(yv), t=float(yv), anom=F.cell(r, 2).value,
                          erf=F.cell(r, 40).value, oni=F.cell(r, 4).value,
                          roni=float(np.mean(rl)) if len(rl) == 12 else None))
    arows = [d for d in arows if isnum(d["roni"])]
    ta = np.array([d["t"] for d in arows])
    sl, it = np.polyfit(ta, np.array([d["roni"] for d in arows], float), 1)
    for d in arows:
        d["roni_dt"] = d["roni"] - (sl * d["t"] + it)
    ya = np.array([d["anom"] for d in arows], float)
    print(f"\n   annual anthropogenic-ERF model (n = {len(arows)})")
    print(f"   {'ENSO control':<24}{'b degC per W/m2':>17}{'b x F2x':>10}"
          f"{'c':>8}{'R^2':>9}{'2023-25 resid':>15}")
    for lab, key in (("ONI v5 (baseline)", "oni"), ("RONI v5", "roni"),
                     ("RONI v5, detrended", "roni_dt")):
        f = ols(ya, [[d["erf"] for d in arows], [d[key] for d in arows]])
        res23 = float(np.mean([ev for d, ev in zip(arows, f["e"]) if d["y"] >= 2023]))
        print(f"   {lab:<24}{f['b'][1]:>17.4f}{f['b'][1]*F2X:>10.4f}"
              f"{f['b'][2]:>8.4f}{f['r2']:>9.4f}{res23:>15.4f}")
    print(f"   workbook baseline: Forcing!T14 b = {F['T14'].value:.4f}, "
          f"T17 b x F2x = {F['T17'].value:.4f}, T31 residual = {F['T31'].value:.4f}")

    # ---------------- B. identification
    print("\nB. Is the GHG-versus-aerosol split identified?")
    def fc(c):
        return np.array([F.cell(r, c).value for r in range(*ANNUAL_ROWS)], float)
    anom_a, ghg, aer, other, oni_a, erf_a = (fc(2), fc(11), fc(9), fc(10),
                                             fc(4), fc(12))
    warm = ghg + other          # Forcing!AQ as corrected in v1.6: K + J
    print(f"   headline regressor L = K + I + J holds to "
          f"{np.abs(erf_a - (ghg + aer + other)).max():.1e}, so the split must use "
          f"K + J against I")
    print(f"   corr(warming terms, aerosol), 1959-2025 = "
          f"{np.corrcoef(warm, aer)[0, 1]:+.4f}   (GHG-only version, "
          f"Forcing!T24, is {F['T24'].value:+.4f})")
    print(f"   aerosol ERF is U-shaped: {aer[0]:.3f} in 1959, {aer.min():.3f} at "
          f"its deepest, {aer[-1]:.3f} in 2025, while the warming terms trend "
          f"{warm[0]:.3f} -> {warm[-1]:.3f} (of which other anthro "
          f"{other[0]:.3f} -> {other[-1]:.3f})")
    print("   for comparison, the superseded GHG-only split:")
    _f0 = ols(anom_a, [ghg, aer, oni_a])
    print(f"      g {_f0['b'][1]:.4f}  h {_f0['b'][2]:.4f}  g-h "
          f"{_f0['b'][1] - _f0['b'][2]:+.4f}   <- omitted-variable artefact")
    f = ols(anom_a, [warm, aer, oni_a])
    print(f"   split regression (R^2 {f['r2']:.4f}, rho {f['rho']:.3f}):")
    for i, nm in enumerate(("intercept", "g, warming-term response",
                            "h, aerosol response", "c, ONI")):
        print(f"      {nm:<22}{f['b'][i]:+.4f}   naive SE {f['se'][i]:.4f}"
              f"   AR(1) SE {f['se'][i]*f['infl']:.4f}")
    _t21, _t22 = F["T21"].value, F["T22"].value
    if isinstance(_t21, (int, float)) and isinstance(_t22, (int, float)):
        print(f"      workbook Forcing!T21 g = {_t21:.4f}, T22 h = {_t22:.4f}")
    else:
        print("      workbook Forcing!T21:T22 carry no cached value (the file "
              "recalculates on open); compare after opening in Excel")
    d = f["b"][1] - f["b"][2]
    sed = math.sqrt(f["V"][1, 1] + f["V"][2, 2] - 2 * f["V"][1, 2]) * f["infl"]
    corr_est = f["V"][1, 2] / math.sqrt(f["V"][1, 1] * f["V"][2, 2])
    print(f"   g - h = {d:+.4f}, AR(1) SE {sed:.4f}, t = {d/sed:+.2f}, "
          f"p = {two_sided_p(d/sed):.3f}")
    print(f"   95% interval on g - h: {d - 1.96*sed:+.3f} to {d + 1.96*sed:+.3f}")
    print(f"   corr(g_hat, h_hat) = {corr_est:+.3f}")
    print("   the data cannot reject g = h, which is what the single-ERF "
          "headline regression assumes")
    print("\n   effect of the aerosol series magnitude (AR6 5-95% is x0.41 to x1.67):")
    print(f"   {'scale s':>10}{'g':>9}{'h':>9}{'h x s':>9}{'R^2':>9}{'t(g-h)':>9}")
    for s in (0.41, 0.60, 1.00, 1.25, 1.67):
        fs = ols(anom_a, [warm, aer * s, oni_a])
        sd2 = math.sqrt(fs["V"][1, 1] + fs["V"][2, 2] - 2 * fs["V"][1, 2]) * fs["infl"]
        print(f"   {s:>10.2f}{fs['b'][1]:>9.4f}{fs['b'][2]:>9.4f}"
              f"{fs['b'][2]*s:>9.4f}{fs['r2']:>9.4f}"
              f"{(fs['b'][1]-fs['b'][2])/sd2:>9.2f}")
    print("   h scales as 1/s while h x s and the R^2 are invariant: the aerosol")
    print("   CONTRIBUTION is identified, the response per W/m2 is not, and the")
    print("   apparent g-versus-h difference is a function of the assumed magnitude.")
    fe = ols(anom_a, [warm + aer, oni_a])
    print(f"   imposing g = h on GHG+aerosol: coefficient {fe['b'][1]:.4f}, "
          f"R^2 {fe['r2']:.4f} against the split's {f['r2']:.4f}")

    # ---------------- C. the albedo candidate
    print("\nC. Can the CERES shortwave decline account for the 2023-25 remainder?")
    yrs, net, lw, sw = [], [], [], []
    for r in range(5, 40):
        yv = C.cell(r, 1).value
        if isinstance(yv, int) and C.cell(r, 2).value == 12:
            yrs.append(yv); net.append(C.cell(r, 3).value)
            lw.append(C.cell(r, 4).value); sw.append(C.cell(r, 5).value)
    yrs = np.array(yrs); net = np.array(net, float)
    lw = np.array(lw, float); sw = np.array(sw, float)
    ref = (yrs >= 2001) & (yrs <= 2022)
    rec = yrs >= 2023
    dsw = sw[rec].mean() - sw[ref].mean()
    dlw = lw[rec].mean() - lw[ref].mean()
    dnet = net[rec].mean() - net[ref].mean()
    b14 = float(F["T14"].value)
    pulse = float(wb["Volcanic"]["N36"].value)
    rem = float(wb["Recent 2023-25"]["D17"].value)
    erf_tot = {int(F.cell(r, 1).value): F.cell(r, 15).value
               for r in range(*ANNUAL_ROWS) if isnum(F.cell(r, 1).value)}
    derf = (np.mean([erf_tot[q] for q in range(2023, 2026)])
            - np.mean([erf_tot[q] for q in range(2001, 2023)]))
    print(f"   CERES annual years {yrs[0]}-{yrs[-1]}; 2023-2025 against 2001-2022:")
    print(f"      reflected SW   {dsw:+.4f} W/m2  (so {-dsw:+.4f} more absorbed)")
    print(f"      outgoing LW    {dlw:+.4f} W/m2")
    print(f"      net imbalance  {dnet:+.4f} W/m2")
    print(f"      total ERF      {derf:+.4f} W/m2")
    print(f"   treating the SW change as a forcing at the sustained response "
          f"b = {b14:.4f} gives {-dsw * b14:+.4f} degC,")
    print(f"   which is {(-dsw * b14) / rem:.1f} times the {rem:.4f} degC remainder "
          f"(Recent 2023-25!D17).")
    print(f"   At the PULSE response {pulse:.4f} (Volcanic!N36), which Forcing!S61 "
          f"argues is the")
    print(f"   right rate for short-lived forcing and which the Recent tab itself "
          f"uses for Hunga")
    print(f"   Tonga over this window, it gives {-dsw * pulse:+.4f} degC, "
          f"{(-dsw * pulse) / rem:.1f} times the remainder.")
    print(f"   {(-dsw) / derf:.2f} times the change in total ERF, and "
          f"{(-dsw) / dnet:.2f} times the change in the")
    print("   CERES net imbalance. A TOA flux change mixes forced cloud and albedo")
    print("   change, feedback to warming and internal variability, and these numbers")
    print("   do not decompose it. What they do establish is that the whole SW change")
    print("   cannot be entered as an independent forcing line: the implied warming")
    print("   exceeds both the residual and the independently estimated forcing change,")
    print("   and the regression's response coefficient already contains whatever part")
    print(f"   of it is feedback. For scale, the aerosol line already in the budget is "
          f"{F['T34'].value:+.4f} degC.")
    print("\n   Conclusion: the leading candidate for the remainder is not usable as a")
    print("   forcing line. The ~0.07 degC stays unexplained, and is still about one")
    print("   standard error of a three-year mean.")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit("usage: open_items.py <workbook.xlsx> <roni_v5_cpc.csv>")
    main(sys.argv[1], sys.argv[2])
